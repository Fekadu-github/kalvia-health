"""PDF and Excel renderers for the period earnings report."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def fee_label(t: str) -> str:
    return t.replace("_", " ").capitalize()


def _period(r: dict) -> str:
    return f"{r['period_start']} to {r['period_end']}"


def _share_rows(r: dict):
    """(provider, fee type, tasks, fees, percent, share) per provider line."""
    rows = []
    for p in r["providers"]:
        if not p["lines"]:
            rows.append((p["full_name"], "No share lines set", p["tasks_count"], p["gross"], None, 0.0))
        for l in p["lines"]:
            rows.append((p["full_name"], fee_label(l["type"]), l["count"], l["fees"], l["percent"], l["share"]))
    return rows


# ---------------- Excel ----------------
def build_xlsx(r: dict) -> bytes:
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE8E2")

    def header(ws, row, titles):
        for i, t in enumerate(titles, 1):
            c = ws.cell(row=row, column=i, value=t)
            c.font = bold
            c.fill = head_fill

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths):
            ws.column_dimensions[chr(65 + i)].width = w

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Earnings report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Period: {_period(r)}"
    summary = [
        ("Cases opened", r["cases_opened"]),
        ("Cases with fees", r["cases_with_fees"]),
        ("Total fees received", r["total_received"]),
        ("Total provider shares", r["total_provider_shares"]),
        ("Platform net", r["platform_net"]),
    ]
    for i, (k, v) in enumerate(summary, 4):
        ws.cell(row=i, column=1, value=k).font = bold
        ws.cell(row=i, column=2, value=v)
    row = 11
    ws.cell(row=row - 1, column=1, value="Fees received by type").font = bold
    header(ws, row, ["Fee type", "Count", "Total"])
    first = row + 1
    for t in r["received_by_type"]:
        row += 1
        ws.cell(row=row, column=1, value=fee_label(t["type"]))
        ws.cell(row=row, column=2, value=t["count"])
        ws.cell(row=row, column=3, value=t["total"])
    row += 1
    ws.cell(row=row, column=1, value="Total").font = bold
    ws.cell(row=row, column=2, value=f"=SUM(B{first}:B{row-1})").font = bold
    ws.cell(row=row, column=3, value=f"=SUM(C{first}:C{row-1})").font = bold
    widths(ws, [28, 14, 16])

    # Provider shares
    ws = wb.create_sheet("Provider shares")
    header(ws, 1, ["Provider", "Fee type", "Tasks", "Fees", "Share %", "Provider share"])
    n = 1
    for prov, fee, count, fees, pct, share in _share_rows(r):
        n += 1
        ws.append([prov, fee, count, fees, pct if pct is not None else "", share])
    ws.cell(row=n + 1, column=1, value="Total").font = bold
    for col in ("C", "D", "F"):
        c = ws[f"{col}{n + 1}"]
        c.value = f"=SUM({col}2:{col}{n})"
        c.font = bold
    widths(ws, [28, 24, 10, 14, 10, 16])

    # Provider totals
    ws = wb.create_sheet("Provider totals")
    header(ws, 1, ["Provider", "Specialty", "Cases worked", "Tasks", "Gross fees", "Shared fees", "Provider share"])
    for p in r["providers"]:
        ws.append([p["full_name"], p["specialty"], p["cases_worked"], p["tasks_count"], p["gross"], p["shared_fees"], p["share_total"]])
    n = len(r["providers"]) + 1
    ws.cell(row=n + 1, column=1, value="Total").font = bold
    for col in ("C", "D", "E", "F", "G"):
        c = ws[f"{col}{n + 1}"]
        c.value = f"=SUM({col}2:{col}{n})"
        c.font = bold
    widths(ws, [28, 22, 14, 10, 14, 14, 16])

    # All fees
    ws = wb.create_sheet("All fees")
    header(ws, 1, ["Date", "Fee type", "Case number", "Provider", "Amount"])
    for f in r["fees"]:
        ws.append([f["date"][:10] if f["date"] else "", fee_label(f["type"]), f["case_number"] or "", f["provider_name"] or "", f["amount"]])
    n = len(r["fees"]) + 1
    ws.cell(row=n + 1, column=1, value="Total").font = bold
    ws[f"E{n + 1}"] = f"=SUM(E2:E{n})"
    ws[f"E{n + 1}"].font = bold
    widths(ws, [14, 22, 20, 28, 14])

    for sheet in wb.worksheets:
        for row_cells in sheet.iter_rows():
            for c in row_cells:
                if isinstance(c.value, float) or (isinstance(c.value, str) and c.value.startswith("=SUM")):
                    c.number_format = "#,##0.00"
        sheet.freeze_panes = "A2" if sheet.title != "Summary" else None
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------- PDF ----------------
def _m(n) -> str:
    return "-" if n is None else f"{n:,.2f}"


def _table(data, col_widths=None, right_cols=(), total_row=False):
    t = Table(data, colWidths=col_widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE8E2")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B8C4BE")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for c in right_cols:
        style.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    if total_row:
        style += [("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F1F5F3"))]
    t.setStyle(TableStyle(style))
    return t


def build_pdf(r: dict) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28, title="Earnings report")
    ss = getSampleStyleSheet()
    story = [
        Paragraph("Earnings report", ss["Title"]),
        Paragraph(f"Period: {_period(r)}", ss["Normal"]),
        Spacer(1, 10),
        _table(
            [["Cases opened", "Cases with fees", "Total fees received", "Provider shares", "Platform net"],
             [r["cases_opened"], r["cases_with_fees"], _m(r["total_received"]), _m(r["total_provider_shares"]), _m(r["platform_net"])]],
            right_cols=(2, 3, 4),
        ),
        Spacer(1, 12),
        Paragraph("Fees received by type", ss["Heading3"]),
    ]
    rows = [["Fee type", "Count", "Total"]] + [[fee_label(t["type"]), t["count"], _m(t["total"])] for t in r["received_by_type"]]
    rows.append(["Total", sum(t["count"] for t in r["received_by_type"]), _m(r["total_received"])])
    story += [_table(rows, col_widths=[220, 80, 120], right_cols=(1, 2), total_row=True), Spacer(1, 12)]

    story.append(Paragraph("Provider shares", ss["Heading3"]))
    sr = _share_rows(r)
    rows = [["Provider", "Fee type", "Tasks", "Fees", "Share %", "Provider share"]]
    for prov, fee, count, fees, pct, share in sr:
        rows.append([prov, fee, count, _m(fees), "-" if pct is None else f"{pct:g}%", _m(share)])
    rows.append(["Total", "", sum(x[2] for x in sr), _m(sum(x[3] for x in sr)), "", _m(sum(x[5] for x in sr))])
    story += [_table(rows, col_widths=[190, 150, 60, 100, 70, 110], right_cols=(2, 3, 4, 5), total_row=True), Spacer(1, 12)]

    story.append(Paragraph("Provider totals", ss["Heading3"]))
    rows = [["Provider", "Cases worked", "Tasks", "Gross fees", "Shared fees", "Provider share"]]
    for p in r["providers"]:
        rows.append([p["full_name"], p["cases_worked"], p["tasks_count"], _m(p["gross"]), _m(p["shared_fees"]), _m(p["share_total"])])
    rows.append(["Total", sum(p["cases_worked"] for p in r["providers"]), sum(p["tasks_count"] for p in r["providers"]),
                 _m(sum(p["gross"] for p in r["providers"])), _m(sum(p["shared_fees"] for p in r["providers"])),
                 _m(sum(p["share_total"] for p in r["providers"]))])
    story += [_table(rows, col_widths=[190, 90, 70, 110, 110, 110], right_cols=(1, 2, 3, 4, 5), total_row=True), Spacer(1, 12)]

    story.append(Paragraph("All fees received", ss["Heading3"]))
    rows = [["Date", "Fee type", "Case number", "Provider", "Amount"]]
    for f in r["fees"]:
        rows.append([f["date"][:10] if f["date"] else "", fee_label(f["type"]), f["case_number"] or "-", f["provider_name"] or "-", _m(f["amount"])])
    rows.append(["Total", "", "", "", _m(r["total_received"])])
    story.append(_table(rows, col_widths=[80, 150, 140, 200, 100], right_cols=(4,), total_row=True))
    doc.build(story)
    return buf.getvalue()
